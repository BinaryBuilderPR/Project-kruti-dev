"""
Comprehensive Deep Corpus Miner & Writing Style Analyzer for '2024 data'.
Scans all DOCX, PDF, Excel, and Text files, extracts Kruti Dev & Unicode text,
mines official department vocabulary, frequent n-gram phrasing, and writing styles.
"""

import os
import sys
import re
import json
from collections import Counter, defaultdict
from typing import List, Dict, Set, Tuple

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
sys.stdout.reconfigure(encoding='utf-8')

import docx
import openpyxl
from pypdf import PdfReader

from krutidev_engine.converter import kruti_to_unicode, unicode_to_kruti


# Validation regex for pure Hindi words
DEV_WORD_PATTERN = re.compile(r'^[\u0900-\u097F]+$')
MATRA_START = re.compile(r'^[\u093E-\u094C\u0901-\u0903\u094D]')
DOUBLE_MATRA = re.compile(r'[\u093E-\u094C]{2,}')


def is_valid_hindi_word(word: str) -> bool:
    """Strictly validates if a string is a clean Hindi word."""
    if len(word) <= 1:
        return False
    if not DEV_WORD_PATTERN.match(word):
        return False
    if MATRA_START.match(word):
        return False
    if DOUBLE_MATRA.search(word):
        return False
    return True


def clean_and_tokenize(text: str) -> List[str]:
    """Cleans text and extracts valid Hindi words."""
    cleaned = re.sub(r'[a-zA-Z0-9\-_.,:;!?"\'\(\)\[\]\/\\%~`@#$^&*+=|0-9०-९“”।]+', ' ', text)
    tokens = [t.strip() for t in cleaned.split() if is_valid_hindi_word(t.strip())]
    return tokens


def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from a DOCX file."""
    text_chunks = []
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            # Check font
            fonts = set(run.font.name for run in p.runs if run.font.name)
            is_kruti = any('kruti' in str(f).lower() for f in fonts)
            if is_kruti:
                text_chunks.append(kruti_to_unicode(t))
            else:
                # If has devanagari, keep as is, otherwise try kruti conversion
                if any('\u0900' <= c <= '\u097F' for c in t):
                    text_chunks.append(t)
                else:
                    text_chunks.append(kruti_to_unicode(t))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    ct = cell.text.strip()
                    if ct:
                        text_chunks.append(kruti_to_unicode(ct))
    except Exception:
        pass
    return "\n".join(text_chunks)


def extract_text_from_pdf(file_path: str, max_pages: int = 15) -> str:
    """Extracts text from a PDF file."""
    text_chunks = []
    try:
        reader = PdfReader(file_path)
        for i, page in enumerate(reader.pages[:max_pages]):
            pt = page.extract_text()
            if pt:
                if any('\u0900' <= c <= '\u097F' for c in pt):
                    text_chunks.append(pt)
                else:
                    text_chunks.append(kruti_to_unicode(pt))
    except Exception:
        pass
    return "\n".join(text_chunks)


def extract_text_from_excel(file_path: str, max_rows: int = 150) -> str:
    """Extracts text from an Excel file."""
    text_chunks = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for sheet in wb.sheetnames[:3]:
            ws = wb[sheet]
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                row_count += 1
                if row_count > max_rows:
                    break
                for val in row:
                    if isinstance(val, str) and val.strip():
                        text_chunks.append(kruti_to_unicode(val.strip()))
        wb.close()
    except Exception:
        pass
    return "\n".join(text_chunks)


def deep_mine_all(data_dir: str = "2024 data") -> Dict:
    print(f"Starting Deep Mining across '{data_dir}'...")
    
    total_files = 0
    docx_count = 0
    pdf_count = 0
    xlsx_count = 0
    
    all_sentences = []
    word_freq = Counter()
    bigram_freq = Counter()
    trigram_freq = Counter()
    
    # Official style extractors
    letter_openings = Counter()
    letter_closings = Counter()
    designations = Counter()
    court_terms = Counter()
    
    for root, dirs, files in os.walk(data_dir):
        for file in files:
            if file.startswith('~$'):
                continue
            ext = os.path.splitext(file)[1].lower()
            fpath = os.path.join(root, file)
            
            extracted_text = ""
            if ext == '.docx':
                extracted_text = extract_text_from_docx(fpath)
                docx_count += 1
            elif ext == '.pdf':
                extracted_text = extract_text_from_pdf(fpath, max_pages=10)
                pdf_count += 1
            elif ext in ('.xlsx', '.xls'):
                extracted_text = extract_text_from_excel(fpath)
                xlsx_count += 1
            elif ext == '.txt':
                try:
                    with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
                        extracted_text = kruti_to_unicode(f.read())
                except:
                    pass
                    
            if not extracted_text:
                continue
                
            total_files += 1
            if total_files % 100 == 0:
                print(f"  Processed {total_files} files (DOCX: {docx_count}, PDF: {pdf_count}, Excel: {xlsx_count})...")
                
            # Process sentences & paragraphs
            lines = [l.strip() for l in extracted_text.split('\n') if l.strip()]
            for line in lines:
                tokens = clean_and_tokenize(line)
                if not tokens:
                    continue
                    
                # Update word counts
                for t in tokens:
                    word_freq[t] += 1
                    
                # Update n-grams
                for i in range(len(tokens) - 1):
                    bigram = f"{tokens[i]} {tokens[i+1]}"
                    bigram_freq[bigram] += 1
                    
                for i in range(len(tokens) - 2):
                    trigram = f"{tokens[i]} {tokens[i+1]} {tokens[i+2]}"
                    trigram_freq[trigram] += 1
                    
                # Identify drafting patterns
                if any(k in line for k in ["आलोक में", "निर्देशानुसार", "अनुरोध है", "सूचित किया जाता है", "आदेश दिया जाता है", "अवलोकन किया जाए"]):
                    letter_openings[line[:80]] += 1
                if any(k in line for k in ["विश्वासभाजन", "भवदीय", "आदेशानुसार", "हस्ताक्षर", "प्रतिलिपि"]):
                    letter_closings[line[:80]] += 1
                if any(k in line for k in ["अधीक्षक", "पदाधिकारी", "निदेशक", "प्रधानाध्यापक", "लिपिक", "सचिव", "उपायुक्त", "शिक्षक"]):
                    designations[line[:60]] += 1
                if any(k in line for k in ["याचिका", "न्यायालय", "तथ्य", "शपथपत्र", "प्रतिवादी", "वादी", "न्यायादेश", "अंतरिम"]):
                    court_terms[line[:70]] += 1

    print(f"\nCompleted Deep Mining:")
    print(f"  Total Processed Files: {total_files}")
    print(f"  Total Unique Hindi Words Discovered: {len(word_freq)}")
    print(f"  Total Unique 2-grams: {len(bigram_freq)}")
    print(f"  Total Unique 3-grams: {len(trigram_freq)}")
    
    # Save Top Clean Words to lexicon_words.txt
    lexicon_txt_path = os.path.join(os.path.dirname(__file__), "..", "word_vsto_csharp", "lexicon_words.txt")
    
    existing_words = set()
    if os.path.exists(lexicon_txt_path):
        with open(lexicon_txt_path, 'r', encoding='utf-8') as f:
            existing_words = set(l.strip() for l in f if l.strip())
            
    # Add words with frequency >= 2 that are strictly valid
    added_new = 0
    for w, freq in word_freq.items():
        if freq >= 2 and is_valid_hindi_word(w) and w not in existing_words:
            existing_words.add(w)
            added_new += 1
            
    with open(lexicon_txt_path, 'w', encoding='utf-8') as f:
        for w in sorted(existing_words):
            f.write(w + '\n')
            
    print(f"  Updated '{lexicon_txt_path}' with {len(existing_words)} total verified words (+{added_new} new words).")
    
    # Generate Writing Style & Insights Report
    insights = {
        "total_files_analyzed": total_files,
        "docx_files": docx_count,
        "pdf_files": pdf_count,
        "excel_files": xlsx_count,
        "total_vocabulary": len(existing_words),
        "top_50_office_words": [w for w, _ in word_freq.most_common(50)],
        "top_25_phrases": [p for p, _ in bigram_freq.most_common(25)],
        "top_20_clauses": [t for t, _ in trigram_freq.most_common(20)],
        "sample_letter_openings": [k for k, _ in letter_openings.most_common(10)],
        "sample_letter_closings": [k for k, _ in letter_closings.most_common(10)],
        "sample_designations": [k for k, _ in designations.most_common(10)],
        "sample_court_patterns": [k for k, _ in court_terms.most_common(10)]
    }
    
    insights_path = os.path.join(os.path.dirname(__file__), "..", "corpus_insights.json")
    with open(insights_path, 'w', encoding='utf-8') as f:
        json.dump(insights, f, ensure_ascii=False, indent=2)
        
    print(f"  Saved full insights to '{insights_path}'.")
    return insights


if __name__ == "__main__":
    deep_mine_all()
